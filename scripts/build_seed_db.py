"""
Build the seed database from the source spreadsheet.

    python scripts/build_seed_db.py

The seed database holds domain data only — customers, products, product usage
and sales policies — and is tracked in git. It is opened read-only at runtime;
everything the agent writes goes to the separate runtime database (see
app/database.py). Rebuild it whenever data.xlsx changes.
"""

from __future__ import annotations

import datetime as dt
import sqlite3
from pathlib import Path
from typing import Any

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SOURCE_XLSX = ROOT / "data.xlsx"
SEED_DB_PATH = ROOT / "data" / "seed.db"

# Sheet name -> DDL. The sheet must carry every column the table declares (under the
# name in COLUMN_SOURCES where the spreadsheet uses an older one); extra sheet columns
# are ignored.
TABLES: dict[str, str] = {
    "customers": """
        CREATE TABLE customers (
            customer_id            TEXT PRIMARY KEY,
            customer_name          TEXT NOT NULL,
            company_stage          TEXT,
            industry               TEXT,
            use_case               TEXT,
            business_model         TEXT,
            country                TEXT,
            annual_payment_volume  INTEGER,
            monthly_transactions   INTEGER,
            average_order_value    REAL,
            primary_pain_point     TEXT,
            fraud_risk_level       TEXT,
            integration_maturity   TEXT,
            created_at             TEXT
        )
    """,
    "stripe_products": """
        CREATE TABLE stripe_products (
            product_id                  TEXT PRIMARY KEY,
            product_name                TEXT NOT NULL,
            product_group               TEXT,
            product_category            TEXT,
            short_description           TEXT,
            primary_use_case            TEXT,
            target_customer_type        TEXT,
            complexity_level            TEXT,
            requires_sales_assist       BOOLEAN,
            requires_compliance_review  BOOLEAN
        )
    """,
    "customer_product_usage": """
        CREATE TABLE customer_product_usage (
            usage_id         TEXT PRIMARY KEY,
            customer_id      TEXT NOT NULL REFERENCES customers(customer_id),
            product_id       TEXT NOT NULL REFERENCES stripe_products(product_id),
            usage_status     TEXT,
            start_date       TEXT,
            monthly_volume   INTEGER,
            monthly_revenue  REAL,
            adoption_level   TEXT,
            notes            TEXT
        )
    """,
    "sales_policy_updates": """
        CREATE TABLE sales_policy_updates (
            policy_id             TEXT PRIMARY KEY,
            product_id            TEXT REFERENCES stripe_products(product_id),
            policy_area           TEXT NOT NULL,
            policy_title          TEXT,
            policy_summary        TEXT,
            handoff_team          TEXT,
            effective_start_date  TEXT,
            effective_end_date    TEXT,
            is_active             BOOLEAN
        )
    """,
}


# table column -> the spreadsheet's column, where the two differ. The spreadsheet predates
# the agent and names the Team a policy hands off to with an older word; the seed column
# uses the glossary's (CONTEXT.md: Handoff, Team).
COLUMN_SOURCES: dict[str, dict[str, str]] = {
    "sales_policy_updates": {"handoff_team": "escalation_team"},
}


def _schema(conn: sqlite3.Connection, table: str) -> dict[str, str]:
    """Column name -> declared type, as SQLite itself sees the created table."""
    return {row[1]: row[2].upper() for row in conn.execute(f"PRAGMA table_info({table})")}


def _normalize(declared_type: str, value: Any) -> Any:
    if value is None or (isinstance(value, str) and value.strip() == ""):
        return None
    if isinstance(value, dt.datetime):
        return value.isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if declared_type == "BOOLEAN":
        if isinstance(value, str):
            return int(value.strip().lower() in {"true", "1", "yes", "y"})
        return int(bool(value))
    return value


def _rows(sheet, schema: dict[str, str]) -> list[tuple]:
    iterator = sheet.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(iterator)]
    sources = COLUMN_SOURCES.get(sheet.title, {})
    wanted = {name: sources.get(name, name) for name in schema}
    missing = [column for column in wanted.values() if column not in header]
    if missing:
        raise ValueError(f"sheet '{sheet.title}' is missing columns: {missing}")
    index = {name: header.index(column) for name, column in wanted.items()}
    rows = []
    for raw in iterator:
        if raw is None or all(v is None for v in raw):
            continue
        rows.append(tuple(_normalize(kind, raw[index[name]]) for name, kind in schema.items()))
    return rows


def build(source: Path = SOURCE_XLSX, target: Path = SEED_DB_PATH) -> dict[str, int]:
    """Rebuild `target` from `source`. Returns row counts per table."""
    if not source.exists():
        raise FileNotFoundError(f"source spreadsheet not found: {source}")
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)

    tmp = target.with_suffix(".building")
    if tmp.exists():
        tmp.unlink()
    conn = sqlite3.connect(tmp)
    counts: dict[str, int] = {}
    try:
        # A rollback journal, not WAL: WAL would leave -wal/-shm files beside a
        # tracked file, and the seed is only ever read at runtime.
        conn.execute("PRAGMA journal_mode=DELETE")
        conn.execute("PRAGMA foreign_keys=ON")
        for table, ddl in TABLES.items():
            conn.execute(ddl)
            schema = _schema(conn, table)
            rows = _rows(workbook[table], schema)
            placeholders = ", ".join("?" for _ in schema)
            conn.executemany(
                f"INSERT INTO {table} ({', '.join(schema)}) VALUES ({placeholders})", rows
            )
            counts[table] = len(rows)
        conn.commit()
        conn.execute("VACUUM")
    finally:
        conn.close()

    tmp.replace(target)  # atomic swap; fails loudly if the target is locked
    return counts


if __name__ == "__main__":
    result = build()
    width = max(len(t) for t in result)
    for table, n in result.items():
        print(f"  {table:<{width}}  {n:>4} rows")
    print(f"seed database written to {SEED_DB_PATH.relative_to(ROOT)}")
